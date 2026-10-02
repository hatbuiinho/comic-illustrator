from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path('/workspace')
BASE = ROOT / 'Tap 53/Part 2'
OUT = ROOT / 'Huong-dan-nguoi-moi-phoi-hop-voi-ChatGPT-de-ve-truyen-tranh.docx'

NAVY = '17324D'; BLUE = '2E74B5'; SKY = 'EAF3F8'; PALE = 'F5F7F9'
GREEN = 'E7F3EC'; GOLD = 'FFF4CE'; RED = 'FCE8E6'; GRAY = '5F6B76'; WHITE = 'FFFFFF'


def run_font(run, size=None, bold=None, color=None, italic=None):
    run.font.name = 'Calibri'
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), 'Calibri')
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), 'Calibri')
    if size is not None: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic
    if color is not None: run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    shd = pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd'); pr.append(shd)
    shd.set(qn('w:fill'), fill)


def margins(cell, top=100, start=140, bottom=100, end=140):
    pr = cell._tc.get_or_add_tcPr()
    node = pr.first_child_found_in('w:tcMar')
    if node is None:
        node = OxmlElement('w:tcMar'); pr.append(node)
    for name, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        x = node.find(qn(f'w:{name}'))
        if x is None: x = OxmlElement(f'w:{name}'); node.append(x)
        x.set(qn('w:w'), str(value)); x.set(qn('w:type'), 'dxa')


def repeat_header(row):
    pr = row._tr.get_or_add_trPr()
    h = OxmlElement('w:tblHeader'); h.set(qn('w:val'), 'true'); pr.append(h)


def table_widths(table, widths, indent=120):
    table.autofit = False
    pr = table._tbl.tblPr
    tw = pr.find(qn('w:tblW'))
    if tw is None: tw = OxmlElement('w:tblW'); pr.append(tw)
    tw.set(qn('w:w'), str(sum(widths))); tw.set(qn('w:type'), 'dxa')
    ti = pr.find(qn('w:tblInd'))
    if ti is None: ti = OxmlElement('w:tblInd'); pr.append(ti)
    ti.set(qn('w:w'), str(indent)); ti.set(qn('w:type'), 'dxa')
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for w in widths:
        g = OxmlElement('w:gridCol'); g.set(qn('w:w'), str(w)); grid.append(g)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            pr = cell._tc.get_or_add_tcPr()
            cw = pr.find(qn('w:tcW'))
            if cw is None: cw = OxmlElement('w:tcW'); pr.append(cw)
            cw.set(qn('w:w'), str(widths[i])); cw.set(qn('w:type'), 'dxa')
            margins(cell); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_field(p):
    p.add_run('Trang ')
    r = p.add_run()
    a = OxmlElement('w:fldChar'); a.set(qn('w:fldCharType'), 'begin')
    b = OxmlElement('w:instrText'); b.set(qn('xml:space'), 'preserve'); b.text = 'PAGE'
    c = OxmlElement('w:fldChar'); c.set(qn('w:fldCharType'), 'end')
    r._r.extend([a, b, c])


def callout(doc, label, text, fill=PALE):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.LEFT
    table_widths(t, [9360]); shade(t.cell(0, 0), fill); repeat_header(t.rows[0])
    p = t.cell(0, 0).paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    r = p.add_run(label + ': '); run_font(r, bold=True, color=NAVY)
    p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def simple_table(doc, headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.LEFT; table_widths(t, widths)
    for i, text in enumerate(headers):
        cell = t.rows[0].cells[i]; shade(cell, SKY)
        p = cell.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text); run_font(r, bold=True, color=NAVY)
    repeat_header(t.rows[0])
    for values in rows:
        cells = t.add_row().cells
        for i, text in enumerate(values):
            cells[i].text = text
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs: run_font(r, size=10)
    table_widths(t, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(3)
    p.add_run(text); return p


def response_box(doc, text):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.LEFT
    table_widths(t, [9360]); shade(t.cell(0, 0), GREEN); repeat_header(t.rows[0])
    p = t.cell(0, 0).paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    r = p.add_run('Bạn có thể trả lời:\n'); run_font(r, bold=True, color=NAVY)
    r = p.add_run(text); run_font(r, italic=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def picture(doc, path, caption, width=6.25):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set('descr', caption); shape._inline.docPr.set('title', caption.split(' — ', 1)[0])
    c = doc.add_paragraph(); c.alignment = WD_ALIGN_PARAGRAPH.CENTER; c.paragraph_format.space_after = Pt(8)
    r = c.add_run(caption); run_font(r, size=9, italic=True, color=GRAY)


def role_block(doc, see, decide, chatgpt):
    simple_table(doc, ['Bạn cần xem', 'Bạn cần quyết định'], [(see, decide)], [4680, 4680])
    callout(doc, 'ChatGPT sẽ tự làm', chatgpt, PALE)


def step_title(doc, n, title, outcome):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10); p.paragraph_format.space_after = Pt(3); p.paragraph_format.keep_with_next = True
    r = p.add_run(f'BƯỚC {n}'); run_font(r, size=10, bold=True, color=BLUE)
    h = doc.add_heading(title, level=1)
    p = doc.add_paragraph(outcome); p.paragraph_format.space_after = Pt(9)
    for r in p.runs: run_font(r, italic=True, color=GRAY)


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5); sec.page_height = Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles['Normal']; normal.font.name = 'Calibri'; normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor.from_string('222222'); normal.paragraph_format.space_after = Pt(6); normal.paragraph_format.line_spacing = 1.2
normal._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri'); normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
for name, size, color, before, after in [('Heading 1', 18, NAVY, 16, 8), ('Heading 2', 14, BLUE, 12, 6), ('Heading 3', 12, NAVY, 9, 4)]:
    s = styles[name]; s.font.name = 'Calibri'; s.font.size = Pt(size); s.font.bold = True; s.font.color.rgb = RGBColor.from_string(color)
    s._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri'); s._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    s.paragraph_format.space_before = Pt(before); s.paragraph_format.space_after = Pt(after); s.paragraph_format.keep_with_next = True
for name in ['List Bullet', 'List Bullet 2']:
    styles[name].font.name = 'Calibri'; styles[name].font.size = Pt(11); styles[name].paragraph_format.space_after = Pt(3)

header = sec.header.paragraphs[0]; header.text = 'HƯỚNG DẪN NGƯỜI MỚI • LÀM VIỆC VỚI CHATGPT'
for r in header.runs: run_font(r, size=8.5, bold=True, color=GRAY)
footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT; add_page_field(footer)
for r in footer.runs: run_font(r, size=9, color=GRAY)

# Cover
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(74); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('HƯỚNG DẪN CHO NGƯỜI MỚI'); run_font(r, size=11, bold=True, color=BLUE)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(8)
r = p.add_run('Cùng ChatGPT làm trọn một page\nminh họa truyện tranh'); run_font(r, size=27, bold=True, color=NAVY)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(24)
r = p.add_run('Bạn cần xem gì • quyết định gì • phản hồi thế nào'); run_font(r, size=13, color=BLUE)
callout(doc, 'Ví dụ thật trong dự án', 'Tập 53 → Part 2 → Chương 5 → page 008. Tài liệu sử dụng chính layout, hai hình minh họa và kết quả kiểm tra đã làm.', GREEN)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before = Pt(28)
r = p.add_run('Không cần biết prompt kỹ thuật • Không cần biết cấu trúc file • Không cần tự chạy kiểm tra'); run_font(r, size=10, bold=True, color=GRAY)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Phiên bản 1.0 • 28/09/2026'); run_font(r, size=9, color=GRAY)
doc.add_page_break()

doc.add_heading('Bạn và ChatGPT chia việc như thế nào?', level=1)
doc.add_paragraph('Bạn là người giữ ý nghĩa câu chuyện và lựa chọn hình ảnh. ChatGPT phụ trách phần đọc hồ sơ, chuẩn bị phương án, tạo ảnh, lưu phiên bản và kiểm tra kỹ thuật.')
simple_table(doc, ['Bạn cần làm', 'ChatGPT tự xử lý'], [
    ('Chọn đúng trang và hình muốn làm.', 'Tìm layout, kịch bản, ghi chú và profile đúng tập.'),
    ('Xác nhận ChatGPT hiểu đúng câu chuyện.', 'Tóm tắt nguồn và chỉ ra chi tiết không được xuất hiện.'),
    ('Chọn cách kể bằng hình.', 'Chuẩn bị 2–3 phương án bố cục dễ so sánh.'),
    ('Nhìn ảnh và nói điểm đúng/sai.', 'Viết mô tả kỹ thuật, tạo ảnh và lưu version mới.'),
    ('Xem ảnh trong trang sách.', 'Kiểm tra vùng chữ, mép cắt, gáy và vùng an toàn.'),
    ('Chọn chính xác bản cuối.', 'Chuyển bản đã chọn vào approved và cập nhật đường dẫn.'),
], [4680, 4680])
callout(doc, 'Bạn không cần học ngay', 'Các từ như manifest, schema, composition contract, profile drift, visual QA hay layout QA. Khi cần, tài liệu sẽ giải thích bằng ngôn ngữ nhìn thấy được.', GOLD)

doc.add_heading('Ví dụ chúng ta sẽ theo dõi', level=2)
doc.add_paragraph('Page 008 có hai hình ở nửa dưới. Hình 1 tạo sự tĩnh lặng: các tỳ-kheo thiền bên hồ sen. Hình 2 chuyển sang niềm hy vọng kín đáo: người lính kỳ cựu đứng xa, chắp tay nhìn vị tỳ-kheo đang giảng pháp.')
picture(doc, BASE / 'tmp/pdfs/page-008-source/spread-08.png', 'Hình 1 — Layout nguồn page 008: chữ ở phía trên, hai vùng minh họa ở phía dưới.')

doc.add_page_break()
step_title(doc, 1, 'Nói rõ bạn muốn làm trang nào', 'Kết quả của bước này: cả hai bên cùng nhìn đúng page và đúng vùng hình.')
role_block(doc,
    'Ảnh layout mà ChatGPT đưa ra; số page, chương, part và vị trí hình trên trang.',
    'Bạn muốn làm toàn page hay một hình cụ thể? Nếu một hình, đó là hình bên trái hay bên phải?',
    'Tự tìm đường dẫn file, mở PDF layout và xác định vùng hình. Bạn không cần tự tra thư mục.')
doc.add_heading('Ví dụ page 008', level=2)
doc.add_paragraph('Trong task thực tế, người dùng chọn làm toàn bộ page 008. Vì vậy workflow phải xử lý đồng thời Hình 1 ở bên trái và Hình 2 ở bên phải, đồng thời giữ nhịp kể giữa hai hình.')
response_box(doc, 'Đúng Tập 53, Part 2, Chương 5, page 008. Tôi muốn làm toàn bộ page 8, gồm cả Hình 1 và Hình 2.')
callout(doc, 'Nếu chưa chắc', 'Bạn chỉ cần nói “Cho tôi xem layout của page 8”. ChatGPT phải đưa đúng hình để bạn xác nhận trước khi đọc sâu.', SKY)

step_title(doc, 2, 'Kiểm tra ChatGPT đã hiểu đúng câu chuyện chưa', 'Kết quả của bước này: có một bản tóm tắt ngắn mà bạn đồng ý trước khi nghĩ đến bố cục.')
role_block(doc,
    'Bản tóm tắt gồm: chuyện gì đang xảy ra, ai xuất hiện, cảm xúc chính, điều gì không được xuất hiện.',
    'Cách hiểu có đúng không? Có thiếu nhân vật, hành động hoặc mối quan hệ quan trọng không?',
    'Đọc notes, shot_notes, kịch bản trước/sau, layout và profile của đúng tập; tách dữ kiện khỏi suy luận.')
doc.add_heading('Bản tóm tắt mẫu của cả hai hình', level=2)
callout(doc, 'Hình 1 — Thiền bên hồ sen', 'Các tỳ-kheo ngồi thiền ban đêm cạnh hồ sen. Trọng tâm là đời sống thanh tịnh, chuyên tâm tu tập và cảm giác yên tĩnh.', GREEN)
callout(doc, 'Hình 2 — Người lính hy vọng', 'Một vị tỳ-kheo giảng cho nhóm cư sĩ; người lính kỳ cựu đứng xa, chắp tay và hướng về vị tỳ-kheo với niềm hy vọng kín đáo.', GREEN)
for text in [
    'Nhịp chung: đời sống thanh tịnh → cộng đồng lắng nghe giáo pháp → niềm hy vọng của người lính.',
    'Hình 1 không biến thành phong cảnh hồ sen vắng người; tính tập thể của nhóm tỳ-kheo phải còn rõ.',
    'Hình 2 không biến thành hai người đứng gần trò chuyện trực tiếp.',
    'Không được xuất hiện sớm: vua Pasenadi, món nợ, tiền kiếp hoặc biểu tượng nghiệp báo của page sau.',
]: bullet(doc, text)
response_box(doc, 'Tôi xác nhận cách hiểu nguồn đúng. Hãy thiết kế composition cho cả Hình 1 và Hình 2.')
doc.add_heading('Nếu ChatGPT hiểu sai', level=2)
response_box(doc, 'Cần sửa cách hiểu: người lính chỉ quan sát từ xa, không bước tới gần và không trò chuyện trực tiếp với vị tỳ-kheo.')

doc.add_page_break()
step_title(doc, 3, 'Chọn cách kể bằng hình', 'Kết quả của bước này: bạn chọn được góc nhìn truyền đúng cảm xúc, chưa cần quan tâm prompt kỹ thuật.')
role_block(doc,
    'Từ 2–3 bản phác hoặc phương án; với mỗi phương án, xem ai được nhìn thấy trước và cảm giác mang lại.',
    'Phương án nào kể đúng nhất? Khoảng cách, ánh mắt và vị trí nhân vật có đúng câu chuyện không?',
    'Thiết kế góc máy, vị trí người, hướng nhìn và vùng an toàn; trình bày bằng hình phác dễ so sánh.')
simple_table(doc, ['Phương án', 'Cách kể cả hai hình', 'Đánh đổi'], [
    ('A. Tĩnh lặng rồi mở ra hy vọng', 'Hình 1 toàn cảnh nhóm thiền; Hình 2 trung-toàn cảnh, người lính đứng xa.', 'Kể chuyện rõ; biểu cảm người lính chỉ đọc ở mức vừa.'),
    ('B. Nhấn mạnh đời sống tập thể', 'Hai hình cùng góc hơi cao, cho thấy nhiều người và không gian rộng.', 'Tính tập thể tốt; nhân vật nhỏ và khó giữ profile rõ.'),
    ('C. Gần gũi, giàu cảm xúc (đã chọn)', 'Hình 1 nhìn qua hoa sen; Hình 2 nhìn qua vai người lính.', 'Cảm xúc mạnh; cần giữ vị giảng sư không bị người lính lấn át.'),
], [2450, 4110, 2800])
callout(doc, 'Cách tự kiểm tra Hình 1', 'Hoa sen chỉ dẫn mắt vào nhóm tỳ-kheo, không được che mất tư thế thiền hoặc biến thành chủ thể duy nhất.', GOLD)
callout(doc, 'Cách tự kiểm tra Hình 2', 'Dù vai người lính ở tiền cảnh, vị tỳ-kheo đang giảng vẫn phải là điểm nhìn chính.', GOLD)
response_box(doc, 'Chọn phương án 3 (C) — gần gũi và giàu cảm xúc. Hình 1 dùng hoa sen tiền cảnh; Hình 2 dùng góc nhìn qua vai người lính.')

step_title(doc, 4, 'Xem ảnh và nói điểm đúng, điểm cần sửa', 'Kết quả của bước này: bạn đưa được phản hồi cụ thể mà không cần biết thuật ngữ hội họa.')
role_block(doc,
    'Hai ảnh được tạo riêng ở kích thước đủ lớn. Xem từng ảnh, sau đó xem chúng như một cặp.',
    'Mỗi ảnh đã đúng nội dung chưa? Hai ảnh có cùng phong cách và nối được cùng một nhịp cảm xúc không?',
    'Tự viết prompt riêng cho từng hình, đọc lại đúng profile trước mỗi lượt, tạo hai output riêng và không làm lẫn nhận dạng.')
picture(doc, BASE / 'pages/page-008/outputs/page008-option3-hinh1-thien-ho-sen-v1.png', 'Hình 2 — Output thật của Hình 1: “Thiền bên hồ sen”, option 3, version 1.')
picture(doc, BASE / 'pages/page-008/outputs/page008-option3-hinh2-nguoi-linh-hy-vong-v1.png', 'Hình 3 — Output thật của Hình 2: “Người lính hy vọng”, option 3, version 1.')

doc.add_heading('Năm câu hỏi đủ dùng cho người mới', level=2)
for text in [
    'Hình 1: có đọc được nhóm tỳ-kheo thiền, hồ sen và sự tĩnh lặng không?',
    'Hình 1: đầu cạo, cà-sa nâu, độ tuổi/vóc dáng và tư thế thiền có tự nhiên không?',
    'Hình 2: vị tỳ-kheo có là focus, còn người lính là khung nhìn và điểm cảm xúc phụ không?',
    'Hình 2: khăn xếp đỏ, tóc/râu, giáp nâu, đôi tay chắp và khoảng cách có đúng không?',
    'Cả hai: có cùng phong cách hoạt hình 2D, không bóng/mờ như 3D và không có lỗi tay chân không?',
]: bullet(doc, text)
response_box(doc, 'Cả hai hình đều đúng nội dung và cảm xúc. Giữ Hình 1 với hoa sen tiền cảnh; giữ Hình 2 với góc nhìn qua vai. Tiếp tục kiểm tra cả hai trong trang sách.')
doc.add_heading('Mẫu phản hồi khi cần sửa', level=2)
response_box(doc, 'Hình 1: giữ nhóm thiền và hồ sen, giảm độ nổi của hoa sen tiền cảnh. Hình 2: giữ góc nhìn qua vai, sửa đôi tay chắp tự nhiên hơn. Tạo version mới cho từng hình cần sửa.')
callout(doc, 'Mẹo phản hồi', 'Luôn nói cả hai phần: “giữ nguyên…” và “sửa…”. Nhờ vậy ChatGPT không vô tình thay đổi phần bạn đã thích.', SKY)

doc.add_page_break()
step_title(doc, 5, 'Xem ảnh trong trang sách', 'Kết quả của bước này: bạn biết ảnh không chỉ đẹp riêng lẻ mà còn hoạt động tốt trong layout thật.')
role_block(doc,
    'Bản xem trước sạch có cả chữ và hai hình. Khi cần, xem thêm bản có đường hướng dẫn vùng an toàn.',
    'Ảnh có che chữ, bị cắt nhân vật, lệch nhịp hoặc quá nổi so với hình bên cạnh không?',
    'Đặt ảnh vào đúng khung, kiểm tỷ lệ, vùng chữ, mép cắt, gáy và các chi tiết quan trọng.')
picture(doc, BASE / 'pages/page-008/qa/page008-option3-v1-clean-preview.png', 'Hình 4 — Bản xem trước sạch của page 008 với cả hai ảnh trong layout thật.')
doc.add_heading('Những gì cần nhìn ở page 008', level=2)
for text in [
    'Hình 1 có tạo được nhịp tĩnh lặng trước khi chuyển sang Hình 2 không?',
    'Hình 2 có tiếp nối tự nhiên từ đời sống tu tập sang niềm hy vọng của người lính không?',
    'Vùng chữ phía trên có còn nguyên và dễ đọc không?',
    'Người lính, đôi tay chắp và vị giảng sư có bị mép cắt hoặc gáy làm mất ý nghĩa không?',
    'Hai hình có cân nhau về độ sáng, màu sắc và mức thu hút không?',
]: bullet(doc, text)
response_box(doc, 'Bố cục trong trang ổn, không che chữ và hai hình nối nhịp tự nhiên. Tôi đồng ý giữ bản này.')

doc.add_heading('Bản có đường kiểm tra dùng để làm gì?', level=2)
picture(doc, BASE / 'pages/page-008/qa/page008-option3-v1-layout-overlay.png', 'Hình 5 — Bản kiểm tra layout: các đường hướng dẫn giúp ChatGPT kiểm vùng an toàn, mép cắt và gáy.')
callout(doc, 'Bạn không cần đọc con số', 'Trong ví dụ thật, hệ thống ghi nhận 21 kiểm tra đạt, không có lỗi/cảnh báo; độ lệch tỷ lệ 1,56% nằm trong dung sai 2%. Người mới chỉ cần xem kết quả đặt trang và xác nhận bằng mắt.', GREEN)

doc.add_page_break()
step_title(doc, 6, 'Chọn đúng bản cuối', 'Kết quả của bước này: ChatGPT biết chính xác file nào được phép trở thành bản chuẩn.')
role_block(doc,
    'Các phiên bản còn là ứng viên, đặt cạnh nhau nếu có nhiều bản.',
    'Bản nào là bản cuối? Có cho phép đưa bản đó vào thư mục approved không?',
    'Đối chiếu kết quả kiểm tra, chuyển đúng file được chọn và cập nhật các đường dẫn liên quan.')
callout(doc, 'Vì sao không chỉ nói “ổn”?', 'Nếu còn nhiều phiên bản, từ “ổn” không cho biết bạn chọn bản nào. Hãy nói đúng tên hoặc số version.', GOLD)
response_box(doc, 'Chọn cả hai bản sau làm bản cuối của page 008 và đưa vào approved:\n1) page008-option3-hinh1-thien-ho-sen-v1.png\n2) page008-option3-hinh2-nguoi-linh-hy-vong-v1.png')
callout(doc, 'Trạng thái thật của ví dụ', 'Cả hai output page 008 đã PASS visual QA và layout QA nhưng production log ghi rõ chưa chuyển vào approved. Cần xác nhận đích danh cả hai file như câu trên.', RED)

doc.add_heading('Toàn bộ cuộc trao đổi mẫu', level=1)
simple_table(doc, ['Lần', 'Bạn nói gì'], [
    ('1', 'Đúng Tập 53, Part 2, Chương 5, page 008. Tôi muốn làm toàn bộ page 8.'),
    ('2', 'Cách hiểu đúng. Hãy thiết kế composition cho cả Hình 1 và Hình 2.'),
    ('3', 'Chọn phương án C: Hình 1 qua hoa sen; Hình 2 qua vai người lính.'),
    ('4', 'Cả hai hình đúng nội dung và cảm xúc; tiếp tục kiểm tra trong trang.'),
    ('5', 'Bố cục trong trang ổn, không che chữ và hai hình nối nhịp tự nhiên.'),
    ('6', 'Chọn đích danh hai filename/version làm bản cuối và đưa vào approved.'),
], [900, 8460])

doc.add_heading('Cách phản hồi để sửa ảnh đúng ý', level=1)
doc.add_paragraph('Một phản hồi tốt không cần dài. Chỉ cần đủ bốn phần: ảnh nào → giữ gì → sửa gì → mức độ thay đổi.')
simple_table(doc, ['Phản hồi chưa rõ', 'Phản hồi dễ thực hiện'], [
    ('“Sửa đẹp hơn.”', '“Giữ bố cục qua vai; sửa đôi tay chắp tự nhiên hơn.”'),
    ('“Nhân vật chưa đúng.”', '“Giữ tư thế; sửa khăn xếp đỏ, giáp nâu và tóc/râu theo profile người lính.”'),
    ('“Ảnh hơi lạ.”', '“Giữ nội dung; giảm độ bóng và cảm giác 3D trên áo giáp.”'),
    ('“Cho gần hơn.”', '“Giữ vị giảng sư trong khung; đưa vai người lính gần hơn khoảng một nhịp.”'),
], [3200, 6160])

doc.add_heading('Công thức phản hồi', level=2)
response_box(doc, 'Hình 1 cần sửa: page008-option3-hinh1-thien-ho-sen-v1.png — giữ nhóm thiền, giảm độ nổi của hoa sen. Hình 2 cần sửa: page008-option3-hinh2-nguoi-linh-hy-vong-v1.png — giữ góc qua vai, sửa đôi tay chắp. Không thay đổi nhịp chung của page.')
callout(doc, 'Nên sửa từng mục tiêu', 'Nếu có nhiều lỗi lớn, hãy ưu tiên lỗi làm sai câu chuyện hoặc sai nhân vật trước. Đừng gom quá nhiều thay đổi mâu thuẫn vào một lượt.', SKY)

doc.add_page_break()
doc.add_heading('Checklist cuối dành cho người duyệt', level=1)
doc.add_paragraph('Bạn có thể dùng trang này mà không cần đọc lại toàn bộ tài liệu.')
for heading, items in [
    ('1. Đúng nguồn', ['☐ Đúng tập, part, chương và page.', '☐ Đúng hình cần làm.', '☐ ChatGPT đã tóm tắt đúng câu chuyện.']),
    ('2. Đúng cách kể', ['☐ Biết người xem nhìn thấy ai trước.', '☐ Khoảng cách và hướng nhìn đúng quan hệ.', '☐ Không đưa nội dung của page sau vào sớm.']),
    ('3. Đúng hình ảnh', ['☐ Nhân vật giống profile.', '☐ Tay, chân, trang phục và tư thế không có lỗi dễ thấy.', '☐ Ảnh giữ phong cách hoạt hình 2D.']),
    ('4. Đúng layout', ['☐ Không che chữ.', '☐ Nhân vật quan trọng không bị cắt.', '☐ Hai hình trên page nối nhịp tự nhiên.']),
    ('5. Đúng bản cuối', ['☐ Đã chỉ rõ filename hoặc version.', '☐ Chỉ cho phép đưa approved sau khi xem bản đặt trang.']),
]:
    doc.add_heading(heading, level=2)
    for item in items: bullet(doc, item)

doc.add_heading('Nếu bạn không biết phải nói gì', level=1)
doc.add_paragraph('Dùng một trong các câu ngắn sau. ChatGPT phải tự tìm phần kỹ thuật còn lại.')
for text in [
    '“Cho tôi xem đúng layout trước khi bắt đầu.”',
    '“Tóm tắt cảnh này bằng ngôn ngữ dễ hiểu.”',
    '“Cho tôi 2–3 cách kể bằng hình và giải thích điểm khác nhau.”',
    '“Tôi chọn phương án này. Hãy tiếp tục tạo ảnh và kiểm tra.”',
    '“Giữ phần đã đúng, chỉ sửa chi tiết tôi nêu.”',
    '“Cho tôi xem ảnh trong layout thật trước khi chọn bản cuối.”',
    '“Liệt kê các phiên bản còn là ứng viên và đề xuất một bản.”',
]: bullet(doc, text)

callout(doc, 'Điều cần nhớ', 'Bạn không cần vận hành workflow. Bạn chỉ cần xác nhận đúng trang, đúng câu chuyện, đúng cách kể, đúng hình và đúng phiên bản cuối. ChatGPT chịu trách nhiệm phần kỹ thuật ở giữa.', GREEN)

doc.core_properties.title = 'Hướng dẫn người mới phối hợp với ChatGPT để làm minh họa truyện tranh'
doc.core_properties.subject = 'Ví dụ thực tế Tập 53, Part 2, Chương 5, page 008'
doc.core_properties.author = 'Comic Illustration Project'
doc.core_properties.keywords = 'người mới, ChatGPT, comic workflow, page 008, hướng dẫn'
doc.save(OUT)
print(OUT)
